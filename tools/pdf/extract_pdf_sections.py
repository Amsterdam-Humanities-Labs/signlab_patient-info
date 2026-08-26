#!/usr/bin/env python3
import sys as _sys
from pathlib import Path as _Path
ROOT = _Path(__file__).resolve().parents[2]   # repository root (/web/hh)
DATA = ROOT / 'data'
_sys.path.insert(0, str(ROOT))                 # so `import db_credentials` works from tools/
"""
Extract Bijwerkingen and Waarschuwingen sections from medication PDFs
and create Word documents with tables.
"""

import os
import re
import fitz  # PyMuPDF
from docx import Document
from docx.shared import Inches, Pt
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

PDF_FOLDER = str(DATA / "pdfs")
OUTPUT_BIJWERKINGEN = str(DATA / "Deel3.docx")
OUTPUT_WAARSCHUWINGEN = str(DATA / "Deel4.docx")


def clean_text_for_xml(text):
    """Remove invalid XML characters from text."""
    import re
    # Remove NULL bytes and control characters (except newline, tab, carriage return)
    # Valid XML chars: #x9 | #xA | #xD | [#x20-#xD7FF] | [#xE000-#xFFFD]
    def valid_xml_char(char):
        codepoint = ord(char)
        return (
            codepoint == 0x9 or
            codepoint == 0xA or
            codepoint == 0xD or
            (0x20 <= codepoint <= 0xD7FF) or
            (0xE000 <= codepoint <= 0xFFFD) or
            (0x10000 <= codepoint <= 0x10FFFF)
        )
    return ''.join(char for char in text if valid_xml_char(char))


def extract_text_from_pdf(pdf_path):
    """Extract all text from a PDF file."""
    try:
        doc = fitz.open(pdf_path)
        text = ""
        for page in doc:
            text += page.get_text()
        doc.close()
        # Clean invalid XML characters
        text = clean_text_for_xml(text)
        return text
    except Exception as e:
        print(f"Error reading {pdf_path}: {e}")
        return ""


def extract_medication_name(text):
    """Extract medication name from 'Samenvatting van [name]'."""
    match = re.search(r'Samenvatting van\s+(\w+)', text, re.IGNORECASE)
    if match:
        return match.group(1)
    return None


def extract_section(text, section_name):
    """Extract a specific section from the text."""
    # Common section headers to use as boundaries
    section_headers = [
        'Bijwerkingen', 'Waarschuwingen', 'Gebruik', 'Waarom gebruikt u dit medicijn',
        'Samenvatting van', 'Uitleg van de apotheker', 'Lees voor gebruik'
    ]

    # Find the section start
    pattern = rf'{section_name}\s*\n'
    match = re.search(pattern, text, re.IGNORECASE)
    if not match:
        # Try without newline requirement
        pattern = rf'{section_name}'
        match = re.search(pattern, text, re.IGNORECASE)
        if not match:
            return ""

    start_pos = match.end()

    # Find the next section header (end boundary)
    end_pos = len(text)
    for header in section_headers:
        if header.lower() == section_name.lower():
            continue
        header_match = re.search(rf'\n{header}', text[start_pos:], re.IGNORECASE)
        if header_match:
            potential_end = start_pos + header_match.start()
            if potential_end < end_pos:
                end_pos = potential_end

    section_text = text[start_pos:end_pos].strip()
    return section_text


def parse_bijwerkingen_to_rows(text):
    """Parse Bijwerkingen section into table rows."""
    rows = []

    if not text:
        return rows

    # Split by "Let op!" markers first
    parts = re.split(r'(Let op!)', text)

    main_content = parts[0] if parts else ""
    let_op_sections = []

    # Reconstruct "Let op!" sections
    i = 1
    while i < len(parts):
        if parts[i] == "Let op!":
            if i + 1 < len(parts):
                let_op_sections.append("Let op!" + parts[i + 1])
                i += 2
            else:
                i += 1
        else:
            i += 1

    # Process main content
    if main_content.strip():
        # Check for "U kunt last krijgen van:" pattern
        krijgen_match = re.search(r'(U kunt last krijgen van[:\s]*)', main_content, re.IGNORECASE)

        if krijgen_match:
            # Everything before "U kunt last krijgen van:"
            before = main_content[:krijgen_match.start()].strip()
            if before:
                for para in split_paragraphs(before):
                    if para.strip():
                        rows.append(para.strip())

            # Find where the bullet list ends
            after_krijgen = main_content[krijgen_match.end():]

            # Collect bullet points
            bullets = []
            lines = after_krijgen.split('\n')
            remaining_lines = []
            in_bullet_list = True

            for line in lines:
                line_stripped = line.strip()
                if not line_stripped:
                    continue
                if line_stripped.startswith('•') or line_stripped.startswith('-') or line_stripped.startswith('·'):
                    if in_bullet_list:
                        bullets.append(line_stripped)
                    else:
                        remaining_lines.append(line_stripped)
                else:
                    # Check if this is a continuation of bullet (like "Daardoor kunt u eerder vallen.")
                    if bullets and in_bullet_list and not any(line_stripped.lower().startswith(x) for x in ['soms', 'heeft u', 'overleg', 'let op']):
                        # This might be a continuation
                        bullets[-1] = bullets[-1] + " " + line_stripped
                    else:
                        in_bullet_list = False
                        remaining_lines.append(line_stripped)

            # Create the "U kunt last krijgen van:" row with bullets
            if bullets:
                bullet_text = "U kunt last krijgen van:\n" + "\n".join(bullets)
                rows.append(bullet_text)

            # Process remaining content as separate paragraphs
            remaining_text = '\n'.join(remaining_lines)
            for para in split_paragraphs(remaining_text):
                if para.strip():
                    rows.append(para.strip())
        else:
            # No "U kunt last krijgen van:" - just split by paragraphs
            for para in split_paragraphs(main_content):
                if para.strip():
                    rows.append(para.strip())

    # Add "Let op!" sections as separate rows
    for let_op in let_op_sections:
        cleaned = clean_let_op_section(let_op)
        if cleaned.strip():
            rows.append(cleaned.strip())

    return rows


def parse_waarschuwingen_to_rows(text):
    """Parse Waarschuwingen section into table rows."""
    rows = []

    if not text:
        return rows

    # Split by "Let op!" markers
    parts = re.split(r'(Let op!)', text)

    # Handle content before first "Let op!"
    if parts and parts[0].strip():
        first_part = parts[0].strip()
        # Check if it starts with something like "U mag autorijden..."
        for para in split_paragraphs(first_part):
            if para.strip():
                rows.append(para.strip())

    # Process each "Let op!" section
    i = 1
    while i < len(parts):
        if parts[i] == "Let op!":
            if i + 1 < len(parts):
                section_text = "Let op!" + parts[i + 1]

                # Check for sub-bullets within this Let op! section
                sub_rows = parse_let_op_with_subbullets(section_text)
                rows.extend(sub_rows)
                i += 2
            else:
                i += 1
        else:
            i += 1

    return rows


def parse_let_op_with_subbullets(text):
    """Parse a Let op! section, splitting sub-bullets into separate rows."""
    rows = []

    lines = text.split('\n')
    current_section = []

    for line in lines:
        stripped = line.strip()
        if not stripped:
            continue

        # Check if this is a sub-bullet (starts with • or -)
        if stripped.startswith('•') or stripped.startswith('-') or stripped.startswith('·'):
            # If we have accumulated content, save it first
            if current_section:
                rows.append('\n'.join(current_section))
                current_section = []
            # Add sub-bullet as its own row
            rows.append(stripped)
        else:
            current_section.append(stripped)

    # Don't forget the last accumulated section
    if current_section:
        rows.append('\n'.join(current_section))

    return rows


def clean_let_op_section(text):
    """Clean up a Let op! section text."""
    # Remove extra whitespace
    text = re.sub(r'\s+', ' ', text)
    return text.strip()


def split_paragraphs(text):
    """Split text into paragraphs."""
    # Split on double newlines or when encountering certain patterns
    paragraphs = []

    # First, normalize newlines
    text = text.replace('\r\n', '\n').replace('\r', '\n')

    # Split by double newlines
    parts = re.split(r'\n\s*\n', text)

    for part in parts:
        part = part.strip()
        if part:
            # Further split if there are sentence boundaries followed by newlines
            sub_parts = re.split(r'(?<=[.!?])\s*\n', part)
            for sub in sub_parts:
                sub = sub.strip()
                if sub:
                    paragraphs.append(sub)

    return paragraphs


def set_cell_margins(cell, top=50, bottom=50, left=100, right=100):
    """Set cell margins in twips (1/20 of a point)."""
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for margin_name, margin_value in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        margin = OxmlElement(f'w:{margin_name}')
        margin.set(qn('w:w'), str(margin_value))
        margin.set(qn('w:type'), 'dxa')
        tcMar.append(margin)
    tcPr.append(tcMar)


def create_document_with_tables(medications_data, output_path, section_type):
    """Create a Word document with tables for each medication."""
    doc = Document()

    first_page = True

    for med_name, rows in sorted(medications_data.items()):
        if not rows:
            continue

        # Add page break (except for first page)
        if not first_page:
            doc.add_page_break()
        first_page = False

        # Add medication name as heading
        heading = doc.add_heading(med_name, level=1)

        # Create table with 2 columns
        table = doc.add_table(rows=1, cols=2)
        table.style = 'Table Grid'

        # Set header row
        header_cells = table.rows[0].cells
        header_cells[0].text = 'Paragraaf'
        header_cells[1].text = 'ID'

        # Make header bold
        for cell in header_cells:
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.bold = True

        # Add data rows
        for row_text in rows:
            row = table.add_row()
            cells = row.cells
            # Clean text for XML compatibility
            cells[0].text = clean_text_for_xml(row_text)
            cells[1].text = ''  # Empty ID column

        # Set column widths
        for row in table.rows:
            row.cells[0].width = Inches(5.5)
            row.cells[1].width = Inches(1.0)

    # Save document
    doc.save(output_path)
    print(f"Created: {output_path}")


def main():
    """Main function to process all PDFs and create Word documents."""

    # Get all PDF files
    pdf_files = sorted([f for f in os.listdir(PDF_FOLDER) if f.endswith('.pdf')])
    print(f"Found {len(pdf_files)} PDF files")

    bijwerkingen_data = {}
    waarschuwingen_data = {}

    for pdf_file in pdf_files:
        pdf_path = os.path.join(PDF_FOLDER, pdf_file)
        print(f"Processing: {pdf_file}")

        # Extract text from PDF
        text = extract_text_from_pdf(pdf_path)
        if not text:
            print(f"  - Could not extract text, skipping")
            continue

        # Get medication name
        med_name = extract_medication_name(text)
        if not med_name:
            print(f"  - Could not find medication name, skipping")
            continue

        print(f"  - Medication: {med_name}")

        # Extract sections
        bijwerkingen_text = extract_section(text, 'Bijwerkingen')
        waarschuwingen_text = extract_section(text, 'Waarschuwingen')

        # Parse into rows
        bijwerkingen_rows = parse_bijwerkingen_to_rows(bijwerkingen_text)
        waarschuwingen_rows = parse_waarschuwingen_to_rows(waarschuwingen_text)

        print(f"  - Bijwerkingen rows: {len(bijwerkingen_rows)}")
        print(f"  - Waarschuwingen rows: {len(waarschuwingen_rows)}")

        # Store data
        if bijwerkingen_rows:
            bijwerkingen_data[med_name] = bijwerkingen_rows
        if waarschuwingen_rows:
            waarschuwingen_data[med_name] = waarschuwingen_rows

    # Create Word documents
    print("\nCreating Word documents...")
    create_document_with_tables(bijwerkingen_data, OUTPUT_BIJWERKINGEN, 'Bijwerkingen')
    create_document_with_tables(waarschuwingen_data, OUTPUT_WAARSCHUWINGEN, 'Waarschuwingen')

    print(f"\nDone! Processed {len(bijwerkingen_data)} medications for Bijwerkingen")
    print(f"Processed {len(waarschuwingen_data)} medications for Waarschuwingen")


if __name__ == '__main__':
    main()
